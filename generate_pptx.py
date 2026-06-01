from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BLUE = RGBColor(0x1E, 0x40, 0xAF)
LIGHT_BLUE = RGBColor(0x3B, 0x82, 0xF6)
ACCENT = RGBColor(0x4F, 0x46, 0xE5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1F, 0x29, 0x37)
GRAY = RGBColor(0x64, 0x74, 0x8B)


def add_blue_bar(slide):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.15), Inches(7.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = BLUE
    shape.line.fill.background()


def add_bottom_bar(slide):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.2), Inches(13.333), Inches(0.3))
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT_BLUE
    shape.line.fill.background()


def set_text(tf, text, size=18, bold=False, color=DARK, alignment=PP_ALIGN.LEFT):
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = alignment
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_bullet(tf, text, size=16, color=DARK, bold=False):
    p = tf.add_paragraph()
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold


# SLIDE 1: Title
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = BLUE

txBox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(11), Inches(2))
tf = txBox.text_frame
tf.word_wrap = True
set_text(tf, "From Zero to Production", size=40, bold=True, color=WHITE)
add_bullet(tf, "Building a Multi-Node AI Agent with RAG,", size=24, color=WHITE)
add_bullet(tf, "Guardrails, and Cloud Deployment", size=24, color=WHITE)

txBox2 = slide.shapes.add_textbox(Inches(1), Inches(5.5), Inches(11), Inches(1))
tf2 = txBox2.text_frame
set_text(tf2, "Chandra Prakash", size=20, color=WHITE)
add_bullet(tf2, "github.com/cprakash0105/research-agent", size=14, color=WHITE)

# SLIDE 2: Problem
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_blue_bar(slide)
add_bottom_bar(slide)

txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(1))
set_text(txBox.text_frame, "The Problem", size=32, bold=True, color=BLUE)

txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11), Inches(5))
tf = txBox.text_frame
tf.word_wrap = True
set_text(tf, 'Most AI demos stop at "call an LLM and print the response"', size=20, color=DARK)
add_bullet(tf, "")
add_bullet(tf, "\u2717  No structured reasoning or planning", size=18, color=GRAY)
add_bullet(tf, "\u2717  No grounding in real sources (hallucination)", size=18, color=GRAY)
add_bullet(tf, "\u2717  No security or governance", size=18, color=GRAY)
add_bullet(tf, "\u2717  No production deployment", size=18, color=GRAY)
add_bullet(tf, "")
add_bullet(tf, "Goal: Build something that thinks in steps, uses tools,", size=18, color=DARK)
add_bullet(tf, "retrieves grounded information, and does it all safely.", size=18, color=DARK)

# SLIDE 3: Solution
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_blue_bar(slide)
add_bottom_bar(slide)

txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(1))
set_text(txBox.text_frame, "The Solution: Research Agent", size=32, bold=True, color=BLUE)

txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11), Inches(5.5))
tf = txBox.text_frame
tf.word_wrap = True
set_text(tf, "An agentic AI assistant that autonomously:", size=20, color=DARK)
add_bullet(tf, "")
add_bullet(tf, "\U0001f9e0  Plans \u2014 Breaks query into targeted sub-questions", size=18, color=DARK)
add_bullet(tf, "\U0001f50d  Researches \u2014 Searches web + indexes uploaded documents", size=18, color=DARK)
add_bullet(tf, "\U0001f4ca  Analyzes \u2014 RAG retrieval for grounded synthesis", size=18, color=DARK)
add_bullet(tf, "\u270d\ufe0f   Writes \u2014 Structured report with citations", size=18, color=DARK)
add_bullet(tf, "\U0001f4ac  Answers \u2014 Follow-up Q&A via RAG", size=18, color=DARK)

# SLIDE 4: Architecture
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_blue_bar(slide)
add_bottom_bar(slide)

txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(1))
set_text(txBox.text_frame, "Agent Architecture (LangGraph State Machine)", size=32, bold=True, color=BLUE)

txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.5), Inches(5.5))
tf = txBox.text_frame
tf.word_wrap = True
set_text(tf, "4-Node Pipeline:", size=20, bold=True, color=DARK)
add_bullet(tf, "")
add_bullet(tf, "User Query", size=16, color=GRAY)
add_bullet(tf, "    \u2193", size=16, color=GRAY)
add_bullet(tf, "\U0001f9e0 PLANNER \u2192 Breaks into 3-5 sub-questions", size=17, color=DARK)
add_bullet(tf, "    \u2193", size=16, color=GRAY)
add_bullet(tf, "\U0001f50d RESEARCHER \u2192 Tavily search + FAISS indexing", size=17, color=DARK)
add_bullet(tf, "    \u2193", size=16, color=GRAY)
add_bullet(tf, "\U0001f4ca ANALYST \u2192 RAG retrieval (top-k) + synthesis", size=17, color=DARK)
add_bullet(tf, "    \u2193", size=16, color=GRAY)
add_bullet(tf, "\u270d\ufe0f  WRITER \u2192 Markdown report with citations", size=17, color=DARK)
add_bullet(tf, "    \u2193", size=16, color=GRAY)
add_bullet(tf, "\U0001f4ac FOLLOW-UP Q&A \u2192 RAG against same index", size=17, color=DARK)

# SLIDE 5: RAG
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_blue_bar(slide)
add_bottom_bar(slide)

txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(1))
set_text(txBox.text_frame, "RAG Pipeline", size=32, bold=True, color=BLUE)

txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11), Inches(5.5))
tf = txBox.text_frame
tf.word_wrap = True
set_text(tf, "Sources (Web + Uploaded Docs)", size=18, bold=True, color=DARK)
add_bullet(tf, "    \u2193 PII Redaction (Microsoft Presidio)", size=16, color=DARK)
add_bullet(tf, "    \u2193 Text Splitting (500 chars, 50 overlap)", size=16, color=DARK)
add_bullet(tf, "    \u2193 Embedding (Gemini Embedding 001, 3072 dim)", size=16, color=DARK)
add_bullet(tf, "    \u2193 FAISS Vector Index (in-memory)", size=16, color=DARK)
add_bullet(tf, "    \u2193 Similarity Search (top-k per query)", size=16, color=DARK)
add_bullet(tf, "    \u2193 Deduplication across sub-questions", size=16, color=DARK)
add_bullet(tf, "    \u2193 Context \u2192 LLM \u2192 Analysis / Answer", size=16, color=DARK)
add_bullet(tf, "")
add_bullet(tf, "Key: 8 relevant chunks > 50 raw search results", size=17, bold=True, color=ACCENT)

# SLIDE 6: Governance
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_blue_bar(slide)
add_bottom_bar(slide)

txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(1))
set_text(txBox.text_frame, "Governance & Security (10 Features)", size=32, bold=True, color=BLUE)

txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(5.5), Inches(5.5))
tf = txBox.text_frame
tf.word_wrap = True
set_text(tf, "Input Protection", size=18, bold=True, color=ACCENT)
add_bullet(tf, "\U0001f512 Authentication (password)", size=15, color=DARK)
add_bullet(tf, "\U0001f6a6 Rate Limiting (10 req/min)", size=15, color=DARK)
add_bullet(tf, "\U0001f6e1\ufe0f  Input Validation (regex)", size=15, color=DARK)
add_bullet(tf, "\U0001f9e0 Prompt Injection Guard (LLM)", size=15, color=DARK)
add_bullet(tf, "\U0001f510 Encrypted Secrets (Fernet)", size=15, color=DARK)

txBox2 = slide.shapes.add_textbox(Inches(6.8), Inches(1.3), Inches(5.5), Inches(5.5))
tf2 = txBox2.text_frame
tf2.word_wrap = True
set_text(tf2, "Output & Data Protection", size=18, bold=True, color=ACCENT)
add_bullet(tf2, "\U0001f6ab Output Safety Filter (LLM)", size=15, color=DARK)
add_bullet(tf2, "\U0001f510 PII Redaction (Presidio)", size=15, color=DARK)
add_bullet(tf2, "\U0001f4cb Audit Trail (JSONL)", size=15, color=DARK)
add_bullet(tf2, "\U0001f4b0 Token Budget Tracking", size=15, color=DARK)
add_bullet(tf2, "\u23f0 Session TTL (30 min)", size=15, color=DARK)

# SLIDE 7: Tech Stack
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_blue_bar(slide)
add_bottom_bar(slide)

txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(1))
set_text(txBox.text_frame, "Technology Stack", size=32, bold=True, color=BLUE)

txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(5.5), Inches(5.5))
tf = txBox.text_frame
tf.word_wrap = True
set_text(tf, "AI & Backend", size=18, bold=True, color=ACCENT)
add_bullet(tf, "LangGraph (agent orchestration)", size=15, color=DARK)
add_bullet(tf, "Google Gemini 2.0 Flash (LLM)", size=15, color=DARK)
add_bullet(tf, "Gemini Embedding 001", size=15, color=DARK)
add_bullet(tf, "Tavily (web search)", size=15, color=DARK)
add_bullet(tf, "FAISS (vector store)", size=15, color=DARK)
add_bullet(tf, "Presidio + spaCy (PII)", size=15, color=DARK)
add_bullet(tf, "Chainlit (UI)", size=15, color=DARK)

txBox2 = slide.shapes.add_textbox(Inches(6.8), Inches(1.3), Inches(5.5), Inches(5.5))
tf2 = txBox2.text_frame
tf2.word_wrap = True
set_text(tf2, "Infrastructure", size=18, bold=True, color=ACCENT)
add_bullet(tf2, "GCP Cloud Run (serverless)", size=15, color=DARK)
add_bullet(tf2, "GCP Secret Manager", size=15, color=DARK)
add_bullet(tf2, "GCP Cloud Build (CI/CD)", size=15, color=DARK)
add_bullet(tf2, "Artifact Registry", size=15, color=DARK)
add_bullet(tf2, "Docker (python:3.11-slim)", size=15, color=DARK)
add_bullet(tf2, "Kind (local K8s dev)", size=15, color=DARK)
add_bullet(tf2, "pytest (76 automated tests)", size=15, color=DARK)

# SLIDE 8: Deployment
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_blue_bar(slide)
add_bottom_bar(slide)

txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(1))
set_text(txBox.text_frame, "Cloud Deployment (GCP Cloud Run)", size=32, bold=True, color=BLUE)

txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11), Inches(5.5))
tf = txBox.text_frame
tf.word_wrap = True
set_text(tf, "CI/CD Pipeline:", size=18, bold=True, color=DARK)
add_bullet(tf, "")
add_bullet(tf, "git push \u2192 Cloud Build \u2192 Artifact Registry \u2192 Cloud Run", size=17, color=ACCENT, bold=True)
add_bullet(tf, "")
add_bullet(tf, "\u2022 Scales to zero when idle (~$0-5/month)", size=16, color=DARK)
add_bullet(tf, "\u2022 Zero-downtime deployments", size=16, color=DARK)
add_bullet(tf, "\u2022 Secrets from GCP Secret Manager", size=16, color=DARK)
add_bullet(tf, "\u2022 Session affinity for WebSocket support", size=16, color=DARK)
add_bullet(tf, "\u2022 1Gi memory, 300s timeout, 0-2 instances", size=16, color=DARK)

# SLIDE 9: Takeaways
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_blue_bar(slide)
add_bottom_bar(slide)

txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(1))
set_text(txBox.text_frame, "Key Takeaways", size=32, bold=True, color=BLUE)

txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11), Inches(5.5))
tf = txBox.text_frame
tf.word_wrap = True
set_text(tf, "1. Agentic AI > prompt chaining", size=17, bold=True, color=DARK)
add_bullet(tf, "   State machines with error handling per node", size=15, color=GRAY)
add_bullet(tf, "")
add_bullet(tf, "2. Governance isn't optional", size=17, color=DARK, bold=True)
add_bullet(tf, "   Rate limiting, PII detection, audit trails = table stakes", size=15, color=GRAY)
add_bullet(tf, "")
add_bullet(tf, "3. RAG quality > RAG quantity", size=17, color=DARK, bold=True)
add_bullet(tf, "   8 relevant chunks beat 50 raw search results", size=15, color=GRAY)
add_bullet(tf, "")
add_bullet(tf, "4. Middleware pattern works for AI governance", size=17, color=DARK, bold=True)
add_bullet(tf, "   Each feature is a layer, toggled independently", size=15, color=GRAY)
add_bullet(tf, "")
add_bullet(tf, "5. Cloud Run = perfect for AI side projects", size=17, color=DARK, bold=True)
add_bullet(tf, "   Scales to zero, WebSocket support, $0 when idle", size=15, color=GRAY)

# SLIDE 10: Resources
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_blue_bar(slide)
add_bottom_bar(slide)

txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(1))
set_text(txBox.text_frame, "Free Learning Resources", size=32, bold=True, color=BLUE)

txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11), Inches(5.5))
tf = txBox.text_frame
tf.word_wrap = True
set_text(tf, "Courses:", size=18, bold=True, color=ACCENT)
add_bullet(tf, "")
add_bullet(tf, "\u2022 Building AI Agents with LangGraph \u2014 DeepLearning.AI", size=15, color=DARK)
add_bullet(tf, "\u2022 AI Agentic Design Patterns \u2014 DeepLearning.AI", size=15, color=DARK)
add_bullet(tf, "\u2022 Building RAG Agents with LLMs \u2014 NVIDIA DLI", size=15, color=DARK)
add_bullet(tf, "\u2022 Generative AI Fundamentals \u2014 IBM CognitiveClass.ai", size=15, color=DARK)
add_bullet(tf, "\u2022 Building GenAI Apps with Python \u2014 IBM CognitiveClass.ai", size=15, color=DARK)
add_bullet(tf, "\u2022 Functions, Tools & Agents \u2014 DeepLearning.AI", size=15, color=DARK)
add_bullet(tf, "")
add_bullet(tf, "Docs:", size=18, color=ACCENT, bold=True)
add_bullet(tf, "\u2022 LangGraph \u2014 langchain-ai.github.io/langgraph", size=15, color=DARK)
add_bullet(tf, "\u2022 Gemini Cookbook \u2014 github.com/google-gemini/cookbook", size=15, color=DARK)

# SLIDE 11: Thank You
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = BLUE

txBox = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(3))
tf = txBox.text_frame
tf.word_wrap = True
set_text(tf, "Thank You", size=44, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
add_bullet(tf, "", size=12, color=WHITE)
add_bullet(tf, "github.com/cprakash0105/research-agent", size=20, color=WHITE)
add_bullet(tf, "research-agent-489654189917.us-central1.run.app", size=18, color=WHITE)
add_bullet(tf, "", size=12, color=WHITE)
add_bullet(tf, "Chandra Prakash", size=20, color=WHITE)

prs.save("Research_Agent_Presentation.pptx")
print("Presentation created successfully!")
